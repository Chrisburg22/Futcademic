#!/usr/bin/env python3
"""
Extrae la lista de asistencia de Gallos San Rafael (Serratos) a JSON.

Entrada : scripts/serratos_src/serratos.xlsx   (descarga de la hoja de Google)
Salida  : scripts/serratos_data.json  + reporte de conteos por hoja en consola

No toca ninguna base de datos. Es el paso de control previo a la carga.

Leyenda de la hoja (la dio la academia):
    A  = Asistencia (verde)          F  = Falta (rojo)
    L  = Lesion o enfermedad         C  = Cordica (azul)
    LG = Liga Gallos                 PP = Partido de preparacion
    Nuevo ingreso (negro)            B1..B100 (cian)

Solo A y F se cargan como asistencia. L, C, LG, PP y los colores sin significado
conocido se CUENTAN en el reporte pero no se escriben: no son ni presencia ni
falta, y adivinarlo sesgaria la etiqueta del modelo.

Uso:  python3 scripts/extract_serratos.py   (requiere openpyxl)
"""
import json
import os
import re
import sys
from collections import Counter
from datetime import date, datetime

import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
XLSX = os.path.join(AQUI, "serratos_src", "serratos.xlsx")
SALIDA = os.path.join(AQUI, "serratos_data.json")
ANIO = 2026

MESES = {"enero": 1, "febrero": 2, "marzo": 3, "abril": 4, "mayo": 5, "junio": 6,
         "julio": 7, "agosto": 8, "septiembre": 9, "setiembre": 9, "octubre": 10,
         "noviembre": 11, "diciembre": 12}
PREFIJO_MES = {k[:5]: v for k, v in MESES.items()}   # SEPTIEM, OCTUBR, AGOST...

VERDE = {"FF00FF00", "FF38761D", "theme7"}
ROJO = {"FFFF0000", "theme5"}


def color(celda):
    f = celda.fill
    if f.fill_type != "solid":
        return None
    fg = f.fgColor
    return fg.rgb if isinstance(fg.rgb, str) else "theme%s" % fg.theme


def clasificar(celda):
    """-> 'A' | 'F' | 'otro' | None. La letra manda; si no hay, manda el color."""
    v = str(celda.value).strip().upper() if celda.value is not None else ""
    if v in ("A", "AA"):
        return "A"
    if v == "F":
        return "F"
    if v:
        return "otro"
    c = color(celda)
    if c in VERDE:
        return "A"
    if c in ROJO:
        return "F"
    return "otro" if c else None


def fecha_nac(v):
    if isinstance(v, (datetime, date)):
        return (v.date() if isinstance(v, datetime) else v).isoformat()
    s = str(v or "").strip().lower()
    m = re.match(r"(\d{1,2})/([a-zñ]+)/(\d{4})$", s)
    if m and m.group(2) in MESES:
        return date(int(m.group(3)), MESES[m.group(2)], int(m.group(1))).isoformat()
    m = re.match(r"(\d{1,2})/(\d{1,2})/(\d{2,4})$", s)
    if m:
        a = int(m.group(3))
        a = a + (2000 if a < 30 else 1900) if a < 100 else a
        try:
            return date(a, int(m.group(2)), int(m.group(1))).isoformat()
        except ValueError:
            return None
    return None        # anio suelto ("2012") o ilegible: no se inventa una fecha


def anio_categoria(nombre):
    """'20-21 A' -> '2020-2021'; 'T/M' -> None."""
    m = re.search(r"(\d{2})-(\d{2})", nombre)
    if m:
        return "20%s-20%s" % m.groups()
    m = re.search(r"(\d{4})-(\d{2})", nombre)
    return "%s-20%s" % m.groups() if m else None


def resolver_dias(dias, mes_final):
    """Los dias cruzan dos meses; el encabezado nombra el ultimo. Se resuelve de atras."""
    fechas, mes = [None] * len(dias), mes_final
    previo = None
    for i in range(len(dias) - 1, -1, -1):
        d = dias[i]
        if d is None:
            continue
        if previo is not None and d > previo:
            mes -= 1
        previo = d
        try:
            fechas[i] = date(ANIO, mes, int(d)).isoformat()
        except ValueError:
            fechas[i] = None
    return fechas


def procesar(ws):
    # fila de encabezado: la que dice "Nombre" en la columna B
    fila_enc = next((r for r in range(1, 12) if str(ws.cell(r, 2).value or "").strip().lower() == "nombre"), None)
    if fila_enc is None:
        return None
    encabezado_mes = " ".join(str(c.value) for r in range(1, fila_enc) for c in ws[r] if isinstance(c.value, str)).lower()
    mes_final = next((n for k, n in PREFIJO_MES.items() if k in encabezado_mes), 10)
    dt = next((str(c.value) for r in range(1, fila_enc) for c in ws[r]
               if isinstance(c.value, str) and c.value.strip().upper().startswith("DT")), None)

    dias, col_dias = [], []
    cols_pago = {}
    for c in ws[fila_enc]:
        v = c.value
        if c.column >= 4 and isinstance(v, (int, float)) and 1 <= v <= 31 and not cols_pago:
            dias.append(int(v)); col_dias.append(c.column)
        elif isinstance(v, str) and c.column >= 4:
            t = v.strip().lower()
            if t.startswith("liga"):
                cols_pago[c.column] = ("liga", None)
            elif t[:5] in PREFIJO_MES:
                cols_pago[c.column] = ("mensualidad", PREFIJO_MES[t[:5]])
    fechas = resolver_dias(dias, mes_final)

    alumnos, vistos = [], set()
    for r in range(fila_enc + 1, ws.max_row + 1):
        nombre = re.sub(r"\s+", " ", str(ws.cell(r, 2).value or "")).strip()
        if len(nombre) < 6 or not re.search(r"[A-Za-zÁÉÍÓÚÑáéíóúñ]{3}", nombre):
            continue
        # la leyenda de colores esta debajo de la lista: "A= Asistencia", "Nuevo Ingreso"...
        if re.match(r"^(A|F|L|C|LG|PP|B\d*|BA)\s*=", nombre, re.I) or nombre.lower() == "nuevo ingreso":
            continue
        if nombre.lower() in vistos:
            continue
        vistos.add(nombre.lower())
        asist, otros = [], Counter()
        for col, f in zip(col_dias, fechas):
            k = clasificar(ws.cell(r, col))
            if k in ("A", "F") and f:
                asist.append({"fecha": f, "present": k == "A"})
            elif k == "otro":
                otros[color(ws.cell(r, col)) or "texto"] += 1
        pagos = []
        for col, (tipo, mes) in cols_pago.items():
            v = ws.cell(r, col).value
            if isinstance(v, (int, float)) and v > 0:
                pagos.append({"tipo": tipo, "mes": mes, "anio": ANIO, "monto": float(v)})
        alumnos.append({"nombre": nombre, "fecha_nac": fecha_nac(ws.cell(r, 3).value),
                        "asistencias": asist, "pagos": pagos, "sin_clasificar": dict(otros)})
    return {"hoja": ws.title, "dt": dt, "anio_categoria": anio_categoria(ws.title),
            "mes_encabezado": mes_final, "fechas": [f for f in fechas if f], "alumnos": alumnos}


def firma(h):
    return tuple((a["nombre"], tuple((x["fecha"], x["present"]) for x in a["asistencias"])) for a in h["alumnos"])


def main():
    if not os.path.exists(XLSX):
        sys.exit("falta %s" % XLSX)
    wb = openpyxl.load_workbook(XLSX)
    hojas, firmas, descartadas = [], {}, []
    for ws in wb.worksheets:
        h = procesar(ws)
        if h is None or not h["alumnos"]:
            descartadas.append((ws.title, "sin alumnos"))
            continue
        f = firma(h)
        if f in firmas:                      # misma hoja repetida (PORTEROS = 2022-23)
            descartadas.append((ws.title, "duplicada de '%s'" % firmas[f]))
            continue
        firmas[f] = ws.title
        hojas.append(h)

    with open(SALIDA, "w", encoding="utf-8") as fh:
        json.dump({"escuela": "Gallos San Rafael Serratos", "hojas": hojas}, fh, ensure_ascii=False, indent=1)

    print("%-14s %7s %9s %6s %7s %7s %-9s" % ("hoja", "alumnos", "marcas", "A", "F", "pagos", "sin clasificar"))
    tot = Counter()
    for h in hojas:
        a = sum(sum(1 for x in al["asistencias"] if x["present"]) for al in h["alumnos"])
        f = sum(sum(1 for x in al["asistencias"] if not x["present"]) for al in h["alumnos"])
        p = sum(len(al["pagos"]) for al in h["alumnos"])
        o = sum(sum(al["sin_clasificar"].values()) for al in h["alumnos"])
        sin_fnac = sum(1 for al in h["alumnos"] if not al["fecha_nac"])
        print("%-14s %7d %9d %6d %7d %7d %5d" % (h["hoja"], len(h["alumnos"]), a + f, a, f, p, o))
        tot.update(alumnos=len(h["alumnos"]), A=a, F=f, pagos=p, otros=o, sin_fnac=sin_fnac)
    print("-" * 62)
    print("%-14s %7d %9d %6d %7d %7d %5d" % ("TOTAL", tot["alumnos"], tot["A"] + tot["F"], tot["A"], tot["F"], tot["pagos"], tot["otros"]))
    print("\nalumnos sin fecha de nacimiento legible: %d" % tot["sin_fnac"])
    for n, why in descartadas:
        print("hoja omitida: %-14s (%s)" % (n, why))
    print("\nescrito: %s" % SALIDA)


if __name__ == "__main__":
    main()
